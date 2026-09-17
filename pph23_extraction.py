"""
Rekap Bukti Potong PPh Unifikasi (BPPU) - PDF ke Excel
========================================================

Cara pakai:
1. Taruh file ini di dalam SATU folder yang sama dengan semua file PDF
   Bukti Potong yang mau direkap.
2. Jalankan (double klik, atau lewat terminal: python rekap_bukpot.py).
3. Program akan otomatis membaca SEMUA file .pdf di folder yang sama,
   lalu membuat file "Rekap_BuktiPotong.xlsx" berisi rekap semua data.

Jika ada file PDF yang gagal dibaca (format beda / rusak / bukan BPPU),
file tersebut akan dilewati dan namanya dicatat di sheet "Gagal" pada
Excel hasil, supaya bisa dicek manual.

Catatan teknis: tabel Kode Objek Pajak/Objek Pajak/DPP/Tarif/PPh (B.3-B.7)
dibaca berdasarkan POSISI KOORDINAT tiap kata di halaman PDF (bukan
menebak urutan baris teks), supaya tetap akurat walau teks "Objek Pajak"
patah ke beberapa baris dengan urutan yang berbeda-beda antar dokumen.
Kalau dalam satu Bukti Potong ada lebih dari satu baris objek pajak,
semuanya akan direkap sebagai baris terpisah di Excel (dengan info
header seperti Nomor Bukti Potong, NPWP, dll tetap sama/diulang).

Requirement: pip install pdfplumber openpyxl
"""

import re
import sys
from pathlib import Path
from datetime import datetime

import pdfplumber
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

# ----------------------------------------------------------------------
# Konfigurasi
# ----------------------------------------------------------------------

OUTPUT_FILENAME = "Rekap_BuktiPotong.xlsx"

HEADERS = [
    "No", "Masa Pajak", "Tahun Pajak", "Pembetulan", "Sifat",
    "Nomor Bukti Potong", "NPWP Penerima", "Nama Penerima", "NITKU Penerima",
    "Jenis Fasilitas", "Jenis PPh", "Kode Objek Pajak", "Objek Pajak",
    "Jumlah Penghasilan Bruto", "Tarif (%)", "PPh Dipotong", "Kode Pasal",
    "Jenis Dokumen", "Tanggal Dokumen", "Nomor Dokumen", "NPWP Pemotong",
    "NITKU Pemotong", "Nama Pemotong", "Tanggal Bukti Potong",
    "Nama Penandatangan",
]

BULAN_ID = {
    "januari": 1, "februari": 2, "maret": 3, "april": 4, "mei": 5,
    "juni": 6, "juli": 7, "agustus": 8, "september": 9,
    "oktober": 10, "november": 11, "desember": 12,
}

KODE_OBJEK_RE = re.compile(r"^\d{2}-\d{3}-\d{2,3}$")
NITKU_RE = re.compile(r"\d{18,26}\s*-\s*[^\n]+")


# ----------------------------------------------------------------------
# Util
# ----------------------------------------------------------------------

def to_number(s):
    """'162.939.049' -> 162939049 (int/float). Kosong -> None."""
    if not s:
        return None
    s = re.sub(r"\s+", "", s).replace(".", "").replace(",", ".")
    try:
        val = float(s)
        return int(val) if val.is_integer() else val
    except ValueError:
        return None


def to_date(s):
    """'24 Juli 2026' -> datetime(2026, 7, 24). Kalau gagal -> teks asli."""
    if not s:
        return None
    s = s.strip()
    m = re.match(r"(\d{1,2})\s+([A-Za-z]+)\s+(\d{4})", s)
    if not m:
        return s
    day, bulan_txt, year = m.groups()
    month = BULAN_ID.get(bulan_txt.lower())
    if not month:
        return s
    try:
        return datetime(int(year), month, int(day))
    except ValueError:
        return s


def get_field(pattern, text, group=1, flags=re.MULTILINE):
    m = re.search(pattern, text, flags)
    return m.group(group).strip() if m else None


def first_line(s):
    """Ambil baris pertama saja dari hasil capture yang mungkin ikut
    menyeret baris label lanjutan di bawahnya."""
    return s.split("\n")[0].strip() if s else s


# ----------------------------------------------------------------------
# Ekstraksi tabel B.3-B.7 berbasis posisi koordinat kata di halaman
# (bukan urutan teks linear), supaya tahan terhadap variasi wrapping.
# ----------------------------------------------------------------------

def extract_table_rows(page):
    words = page.extract_words()
    header = {w["text"]: w for w in words if w["text"] in ("B.3", "B.4", "B.5", "B.6", "B.7")}
    if len(header) < 5:
        return None

    x_dpp = header["B.5"]["x0"]
    x_tarif = header["B.6"]["x0"]
    x_pph = header["B.7"]["x0"]
    header_bottom = max(w["bottom"] for w in header.values())

    b8_candidates = [w for w in words if w["text"] == "B.8" and w["top"] > header_bottom and w["x0"] < 40]
    if not b8_candidates:
        return None
    b8_top = min(w["top"] for w in b8_candidates)

    table_words = [w for w in words if header_bottom < w["top"] < b8_top]
    if not table_words:
        return None

    def col_of(w):
        if KODE_OBJEK_RE.match(w["text"]):
            return "kode"
        x = w["x0"]
        if x < x_dpp - 5:
            return "objek"
        if x < x_tarif - 5:
            return "dpp"
        if x < x_pph - 5:
            return "tarif"
        return "pph"

    for w in table_words:
        w["_col"] = col_of(w)

    row_starts = sorted(set(round(w["top"], 1) for w in table_words if w["_col"] == "kode"))
    if not row_starts:
        return None

    def join_col(row_words, col, multiline=False):
        ws = [w for w in row_words if w["_col"] == col]
        if not ws:
            return ""
        lines = {}
        for w in ws:
            lines.setdefault(round(w["top"], 1), []).append(w)
        out = [
            " ".join(w["text"] for w in sorted(lines[t], key=lambda w: w["x0"]))
            for t in sorted(lines)
        ]
        return "\n".join(out) if multiline else " ".join(out)

    rows = []
    for i, start_top in enumerate(row_starts):
        end_top = row_starts[i + 1] if i + 1 < len(row_starts) else b8_top
        row_words = [w for w in table_words if start_top - 0.5 <= w["top"] < end_top - 0.5]
        rows.append({
            "kode": join_col(row_words, "kode"),
            "objek": join_col(row_words, "objek", multiline=True),
            "dpp": join_col(row_words, "dpp"),
            "tarif": join_col(row_words, "tarif"),
            "pph": join_col(row_words, "pph"),
        })
    return rows


# ----------------------------------------------------------------------
# Ekstraksi satu file PDF -> list of dict (bisa >1 kalau objek pajak
# lebih dari satu baris)
# ----------------------------------------------------------------------

def extract_bukpot(pdf_path: Path) -> list:
    with pdfplumber.open(pdf_path) as pdf:
        page = pdf.pages[0]
        text = page.extract_text() or ""

        if "BUKTI PEMOTONGAN DAN/ATAU PEMUNGUTAN PPH" not in text.upper():
            raise ValueError("Bukan format Bukti Potong PPh Unifikasi (BPPU) yang dikenali.")

        base = {}

        # --- Header: Nomor, Masa Pajak, Sifat, Status/Pembetulan ---
        m = re.search(
            r"^(\S+)\s+(\d{2}-\d{4})\s+(TIDAK\s+FINAL|FINAL)\s+(.+)$",
            text, re.MULTILINE,
        )
        if not m:
            raise ValueError("Baris header NOMOR/MASA PAJAK/SIFAT tidak ditemukan.")
        nomor_bukpot, masa_pajak, sifat, status = m.groups()
        base["Nomor Bukti Potong"] = nomor_bukpot.strip()
        base["Masa Pajak"] = masa_pajak.strip()
        base["Tahun Pajak"] = masa_pajak.strip().split("-")[-1]
        base["Sifat"] = re.sub(r"\s+", " ", sifat.strip())
        base["Pembetulan"] = status.strip()

        # --- A. Penerima Penghasilan ---
        base["NPWP Penerima"] = get_field(r"A\.1\s+NPWP\s*/\s*NIK\s*:\s*(\S+)", text)
        base["Nama Penerima"] = first_line(get_field(r"A\.2\s+NAMA\s*:\s*(.+)", text, flags=re.MULTILINE | re.DOTALL))

        # --- B.1, B.2 ---
        base["Jenis Fasilitas"] = first_line(get_field(r"B\.1\s+Jenis Fasilitas\s*:\s*(.+)", text, flags=re.MULTILINE | re.DOTALL))
        base["Jenis PPh"] = first_line(get_field(r"B\.2\s+Jenis PPh\s*:\s*(.+)", text, flags=re.MULTILINE | re.DOTALL))

        # --- NITKU Penerima (A.3) & NITKU Pemotong (C.2) ---
        # Dicari lewat pola nilai yang khas (nomor panjang + " - " + nama),
        # bukan lewat urutan baris label, karena urutan baris label vs nilai
        # ternyata bisa berbeda-beda antar dokumen (label kadang wrap SEBELUM
        # nilai, kadang SESUDAH nilai).
        nitku_matches = NITKU_RE.findall(text)
        base["NITKU Penerima"] = nitku_matches[0].strip() if len(nitku_matches) >= 1 else None
        base["NITKU Pemotong"] = nitku_matches[1].strip() if len(nitku_matches) >= 2 else None

        # --- B.8, B.9 ---
        jenis_dok_match = re.search(r"Jenis Dokumen\s*:\s*(.+?)\s+Tanggal\s*:\s*(.+)", text)
        if jenis_dok_match:
            base["Jenis Dokumen"] = first_line(jenis_dok_match.group(1))
            base["Tanggal Dokumen"] = to_date(first_line(jenis_dok_match.group(2)))
        else:
            base["Jenis Dokumen"] = None
            base["Tanggal Dokumen"] = None
        base["Nomor Dokumen"] = get_field(r"B\.9\s+Nomor Dokumen\s*:\s*(.+)", text)

        # --- C. Pemotong/Pemungut ---
        base["NPWP Pemotong"] = get_field(r"C\.1\s+NPWP\s*/\s*NIK\s*:\s*(\S+)", text)
        base["Nama Pemotong"] = first_line(get_field(r"C\.3.*?PPh\s*:\s*(.+)", text, flags=re.MULTILINE | re.DOTALL)) \
            or first_line(get_field(r"C\.3.*?:\s*(.+)", text, flags=re.MULTILINE | re.DOTALL))
        base["Tanggal Bukti Potong"] = to_date(get_field(r"C\.4\s+TANGGAL\s*:\s*(.+)", text))
        base["Nama Penandatangan"] = get_field(r"C\.5\s+NAMA PENANDATANGAN\s*:\s*(.+)", text)

        # --- Tabel B.3-B.7 (berbasis posisi koordinat) ---
        table_rows = extract_table_rows(page)
        if not table_rows:
            raise ValueError("Tabel Kode Objek Pajak / DPP / Tarif / PPh (B.3-B.7) tidak ditemukan.")

    results = []
    for t in table_rows:
        row = dict(base)
        row["Kode Objek Pajak"] = t["kode"].strip()
        row["Objek Pajak"] = t["objek"].strip()
        row["Jumlah Penghasilan Bruto"] = to_number(t["dpp"])
        row["Tarif (%)"] = to_number(t["tarif"])
        row["PPh Dipotong"] = to_number(t["pph"])
        row["Kode Pasal"] = ""  # tidak ada di format BPPU ini
        results.append(row)
    return results


# ----------------------------------------------------------------------
# Main
# ----------------------------------------------------------------------

def main():
    if getattr(sys, "frozen", False):
        # Dijalankan sebagai .exe hasil PyInstaller -> pakai lokasi file .exe itu sendiri,
        # BUKAN sys._MEIPASS (itu folder sementara ekstraksi PyInstaller, selalu kosong).
        folder = Path(sys.executable).resolve().parent
    else:
        # Dijalankan sebagai script .py biasa
        folder = Path(__file__).resolve().parent

    pdf_files = sorted(folder.glob("*.pdf"))

    if not pdf_files:
        print(f"Tidak ada file PDF ditemukan di folder ini ({folder}).")
        input("Tekan ENTER untuk keluar...")
        return

    print(f"Ditemukan {len(pdf_files)} file PDF. Memproses...\n")

    rows = []
    failed = []

    for pdf_path in pdf_files:
        try:
            data_list = extract_bukpot(pdf_path)
            rows.extend(data_list)
            suffix = f" ({len(data_list)} baris objek pajak)" if len(data_list) > 1 else ""
            print(f"  [OK]    {pdf_path.name}{suffix}")
        except Exception as e:
            failed.append((pdf_path.name, str(e)))
            print(f"  [GAGAL] {pdf_path.name} -> {e}")

    # ------------------------------------------------------------
    # Tulis Excel
    # ------------------------------------------------------------
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="305496")
    ws.append(HEADERS)
    for col_idx in range(1, len(HEADERS) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"

    for no, data in enumerate(rows, start=1):
        row_values = [no] + [data.get(h, "") for h in HEADERS[1:]]
        ws.append(row_values)

    # lebar kolom otomatis (sederhana)
    for col_idx, header in enumerate(HEADERS, start=1):
        letter = get_column_letter(col_idx)
        max_len = max(len(header), 12)
        ws.column_dimensions[letter].width = min(max_len + 4, 40)

    # format tanggal & angka
    date_cols = [HEADERS.index("Tanggal Dokumen") + 1, HEADERS.index("Tanggal Bukti Potong") + 1]
    num_cols = [HEADERS.index("Jumlah Penghasilan Bruto") + 1, HEADERS.index("PPh Dipotong") + 1]
    for r in range(2, ws.max_row + 1):
        for c in date_cols:
            cell = ws.cell(row=r, column=c)
            if isinstance(cell.value, datetime):
                cell.number_format = "DD-MMM-YYYY"
        for c in num_cols:
            cell = ws.cell(row=r, column=c)
            if isinstance(cell.value, (int, float)):
                cell.number_format = "#,##0"

    # sheet daftar gagal (kalau ada)
    if failed:
        ws2 = wb.create_sheet("Gagal")
        ws2.append(["Nama File", "Alasan"])
        for cell in ws2[1]:
            cell.font = header_font
            cell.fill = header_fill
        for name, reason in failed:
            ws2.append([name, reason])
        ws2.column_dimensions["A"].width = 45
        ws2.column_dimensions["B"].width = 60

    out_path = folder / OUTPUT_FILENAME
    wb.save(out_path)

    print(f"\nSelesai. {len(rows)} baris data, {len(failed)} file gagal.")
    print(f"Hasil disimpan di: {out_path}")
    input("\nTekan ENTER untuk keluar...")


if __name__ == "__main__":
    main()
